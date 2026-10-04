"""
AuraCare Hospital Management System - Backend REST API
Database: hospital_database.xlsx (Multi-Sheet Excel Workbook)
Engine: Python 3, Flask, openpyxl, Flask-CORS
"""

import os
from datetime import datetime
from flask import Flask, request, jsonify
from flask_cors import CORS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
# Enable CORS for all routes so the React frontend can communicate without browser blocking
CORS(app, resources={r"/api/*": {"origins": "*"}})

# Path to the Excel workbook database
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "hospital_database.xlsx")

# Schema definitions and required column headers for each worksheet
SCHEMAS = {
    "Patients": [
        "id", "fullName", "age", "gender", "phone", "email",
        "bloodGroup", "address", "emergencyContact", "department",
        "assignedDoctor", "registrationDate", "status", "notes"
    ],
    "Doctors": [
        "id", "name", "specialization", "department", "phone", "email",
        "experience", "availability", "consultationFee", "status",
        "room", "qualification"
    ],
    "Appointments": [
        "id", "patientId", "patientName", "doctorId", "doctorName",
        "department", "appointmentDate", "appointmentTime", "reason",
        "status", "tokenNumber"
    ],
    "Admissions": [
        "id", "patientId", "patientName", "roomNumber", "bedNumber",
        "department", "doctor", "admissionDate", "expectedDischarge",
        "actualDischarge", "status", "wardType"
    ],
    "Medicines": [
        "id", "name", "category", "quantity", "unitPrice", "expiryDate",
        "supplier", "stockStatus", "dosage", "batchNumber"
    ],
    "Bills": [
        "id", "patientId", "patientName", "consultationFee",
        "medicineCharges", "roomCharges", "otherCharges", "totalAmount",
        "paymentStatus", "billingDate", "paymentMethod", "insuranceProvider"
    ]
}


def init_excel_database():
    """Initializes the Excel workbook if it does not already exist."""
    if os.path.exists(EXCEL_FILE):
        return

    wb = openpyxl.Workbook()
    # Remove default sheet created by openpyxl
    default_sheet = wb.active

    for i, (sheet_name, columns) in enumerate(SCHEMAS.items()):
        ws = wb.create_sheet(title=sheet_name)
        # Append header row
        ws.append(columns)

        # Style header row with teal banner
        header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    wb.remove(default_sheet)
    wb.save(EXCEL_FILE)
    print(f"[AuraCare DBMS] Initialized new workbook at {EXCEL_FILE}")


# Ensure workbook exists on application launch
init_excel_database()


# -------------------------------------------------------------------------
# EXCEL READ / WRITE / UPDATE / DELETE HELPERS
# -------------------------------------------------------------------------

def load_sheet_records(sheet_name):
    """Reads all rows from an Excel worksheet and returns a list of dictionaries."""
    if not os.path.exists(EXCEL_FILE):
        init_excel_database()

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return []

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) <= 1:
        return []

    # First row is headers
    headers = [str(h).strip() for h in rows[0] if h is not None]
    records = []

    for row_values in rows[1:]:
        # Skip completely empty rows
        if not any(row_values):
            continue
        record = {}
        for idx, header in enumerate(headers):
            val = row_values[idx] if idx < len(row_values) else None
            # Normalize None to empty string or appropriate types
            record[header] = "" if val is None else val
        records.append(record)

    return records


def append_sheet_record(sheet_name, record_dict):
    """Appends a new row to the specified Excel worksheet preserving existing data."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(SCHEMAS[sheet_name])
    else:
        ws = wb[sheet_name]

    headers = SCHEMAS[sheet_name]
    row_data = [record_dict.get(h, "") for h in headers]
    ws.append(row_data)

    wb.save(EXCEL_FILE)
    wb.close()
    return record_dict


def update_sheet_record(sheet_name, record_id, updated_fields):
    """Finds a matching ID in column 1, updates the row, and saves the workbook."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return None

    ws = wb[sheet_name]
    headers = SCHEMAS[sheet_name]

    matching_row_idx = None
    for row_idx in range(2, ws.max_row + 1):
        cell_id = ws.cell(row=row_idx, column=1).value
        if cell_id is not None and str(cell_id).strip() == str(record_id).strip():
            matching_row_idx = row_idx
            break

    if matching_row_idx is None:
        wb.close()
        return None

    # Update columns
    for col_idx, header in enumerate(headers, start=1):
        if header in updated_fields:
            ws.cell(row=matching_row_idx, column=col_idx, value=updated_fields[header])

    wb.save(EXCEL_FILE)
    wb.close()

    # Return refreshed updated record
    records = load_sheet_records(sheet_name)
    for r in records:
        if str(r.get("id")).strip() == str(record_id).strip():
            return r
    return updated_fields


def delete_sheet_record(sheet_name, record_id):
    """Deletes the row matching record_id from the specified sheet."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return False

    ws = wb[sheet_name]
    matching_row_idx = None
    for row_idx in range(2, ws.max_row + 1):
        cell_id = ws.cell(row=row_idx, column=1).value
        if cell_id is not None and str(cell_id).strip() == str(record_id).strip():
            matching_row_idx = row_idx
            break

    if matching_row_idx is None:
        wb.close()
        return False

    ws.delete_rows(matching_row_idx, 1)
    wb.save(EXCEL_FILE)
    wb.close()
    return True


def generate_unique_id(sheet_name, prefix, start_number):
    """Generates an incremental unique ID based on existing rows."""
    records = load_sheet_records(sheet_name)
    max_id = start_number
    for r in records:
        val = str(r.get("id", ""))
        if val.startswith(prefix):
            try:
                num = int(val.replace(prefix, ""))
                if num > max_id:
                    max_id = num
            except ValueError:
                pass
    return f"{prefix}{max_id + 1}"


# -------------------------------------------------------------------------
# ROOT & HEALTH CHECK
# -------------------------------------------------------------------------
@app.route("/", methods=["GET"])
def health_check():
    return jsonify({
        "status": "online",
        "service": "AuraCare Hospital DBMS REST API",
        "database": "hospital_database.xlsx",
        "timestamp": datetime.now().isoformat(),
        "sheets": list(SCHEMAS.keys())
    }), 200


# =========================================================================
# 1. PATIENTS REST API ENDPOINTS
# =========================================================================

@app.route("/api/patients", methods=["GET"])
def get_patients():
    try:
        patients = load_sheet_records("Patients")
        return jsonify(patients), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Patients sheet", "details": str(e)}), 500


@app.route("/api/patients", methods=["POST"])
def add_patient():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    # Validate required fields
    if not data.get("fullName"):
        return jsonify({"error": "Field 'fullName' is required"}), 400

    # Generate Unique Patient ID if not provided
    patient_id = data.get("id") or generate_unique_id("Patients", "PAT-", 1000)

    new_patient = {
        "id": patient_id,
        "fullName": data.get("fullName", "").strip(),
        "age": int(data.get("age", 30)) if str(data.get("age", "")).isdigit() else data.get("age", 30),
        "gender": data.get("gender", "Male"),
        "phone": data.get("phone", ""),
        "email": data.get("email", ""),
        "bloodGroup": data.get("bloodGroup", "O+"),
        "address": data.get("address", ""),
        "emergencyContact": data.get("emergencyContact", ""),
        "department": data.get("department", "General Medicine"),
        "assignedDoctor": data.get("assignedDoctor", "Dr. Julian Hayes"),
        "registrationDate": data.get("registrationDate") or datetime.now().strftime("%Y-%m-%d"),
        "status": data.get("status", "Active"),
        "notes": data.get("notes", "")
    }

    try:
        created = append_sheet_record("Patients", new_patient)
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to save to Patients sheet", "details": str(e)}), 500


@app.route("/api/patients/<patient_id>", methods=["PUT"])
def update_patient_route(patient_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    try:
        updated = update_sheet_record("Patients", patient_id, data)
        if not updated:
            return jsonify({"error": f"Patient with ID '{patient_id}' not found"}), 404
        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update patient", "details": str(e)}), 500


@app.route("/api/patients/<patient_id>", methods=["DELETE"])
def delete_patient_route(patient_id):
    try:
        success = delete_sheet_record("Patients", patient_id)
        if not success:
            return jsonify({"error": f"Patient with ID '{patient_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Patient '{patient_id}' removed from Excel sheet", "id": patient_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete patient", "details": str(e)}), 500


# =========================================================================
# 2. DOCTORS REST API ENDPOINTS
# =========================================================================

@app.route("/api/doctors", methods=["GET"])
def get_doctors():
    try:
        doctors = load_sheet_records("Doctors")
        return jsonify(doctors), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Doctors sheet", "details": str(e)}), 500


@app.route("/api/doctors", methods=["POST"])
def add_doctor():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    if not data.get("name"):
        return jsonify({"error": "Field 'name' is required"}), 400

    doctor_id = data.get("id") or generate_unique_id("Doctors", "DOC-", 200)

    new_doctor = {
        "id": doctor_id,
        "name": data.get("name", "").strip(),
        "specialization": data.get("specialization", "General Medicine"),
        "department": data.get("department", "General Medicine"),
        "phone": data.get("phone", ""),
        "email": data.get("email", ""),
        "experience": data.get("experience", "5 Years"),
        "availability": data.get("availability", "Available"),
        "consultationFee": float(data.get("consultationFee", 100)),
        "status": data.get("status", "Active"),
        "room": data.get("room", "Room 101"),
        "qualification": data.get("qualification", "MD")
    }

    try:
        created = append_sheet_record("Doctors", new_doctor)
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to save to Doctors sheet", "details": str(e)}), 500


@app.route("/api/doctors/<doctor_id>", methods=["PUT"])
def update_doctor_route(doctor_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    try:
        updated = update_sheet_record("Doctors", doctor_id, data)
        if not updated:
            return jsonify({"error": f"Doctor with ID '{doctor_id}' not found"}), 404
        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update doctor", "details": str(e)}), 500


@app.route("/api/doctors/<doctor_id>", methods=["DELETE"])
def delete_doctor_route(doctor_id):
    try:
        success = delete_sheet_record("Doctors", doctor_id)
        if not success:
            return jsonify({"error": f"Doctor with ID '{doctor_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Doctor '{doctor_id}' removed from Excel sheet", "id": doctor_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete doctor", "details": str(e)}), 500


# =========================================================================
# 3. APPOINTMENTS REST API ENDPOINTS
# =========================================================================

@app.route("/api/appointments", methods=["GET"])
def get_appointments():
    try:
        appointments = load_sheet_records("Appointments")
        return jsonify(appointments), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Appointments sheet", "details": str(e)}), 500


@app.route("/api/appointments", methods=["POST"])
def add_appointment():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    if not data.get("patientName") or not data.get("doctorName"):
        return jsonify({"error": "Fields 'patientName' and 'doctorName' are required"}), 400

    apt_id = data.get("id") or generate_unique_id("Appointments", "APT-", 3000)
    existing = load_sheet_records("Appointments")

    new_apt = {
        "id": apt_id,
        "patientId": data.get("patientId", "PAT-1001"),
        "patientName": data.get("patientName", ""),
        "doctorId": data.get("doctorId", "DOC-201"),
        "doctorName": data.get("doctorName", ""),
        "department": data.get("department", "General Medicine"),
        "appointmentDate": data.get("appointmentDate") or datetime.now().strftime("%Y-%m-%d"),
        "appointmentTime": data.get("appointmentTime", "10:00 AM"),
        "reason": data.get("reason", "Consultation"),
        "status": data.get("status", "Scheduled"),
        "tokenNumber": data.get("tokenNumber") or (len(existing) + 1)
    }

    try:
        created = append_sheet_record("Appointments", new_apt)
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to save appointment", "details": str(e)}), 500


@app.route("/api/appointments/<apt_id>", methods=["PUT"])
def update_appointment_route(apt_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    try:
        updated = update_sheet_record("Appointments", apt_id, data)
        if not updated:
            return jsonify({"error": f"Appointment with ID '{apt_id}' not found"}), 404
        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update appointment", "details": str(e)}), 500


@app.route("/api/appointments/<apt_id>", methods=["DELETE"])
def delete_appointment_route(apt_id):
    try:
        success = delete_sheet_record("Appointments", apt_id)
        if not success:
            return jsonify({"error": f"Appointment with ID '{apt_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Appointment '{apt_id}' deleted", "id": apt_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete appointment", "details": str(e)}), 500


# =========================================================================
# 4. ADMISSIONS REST API ENDPOINTS
# =========================================================================

@app.route("/api/admissions", methods=["GET"])
def get_admissions():
    try:
        admissions = load_sheet_records("Admissions")
        return jsonify(admissions), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Admissions sheet", "details": str(e)}), 500


@app.route("/api/admissions", methods=["POST"])
def add_admission():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    if not data.get("patientName") or not data.get("roomNumber"):
        return jsonify({"error": "Fields 'patientName' and 'roomNumber' are required"}), 400

    adm_id = data.get("id") or generate_unique_id("Admissions", "ADM-", 400)

    new_adm = {
        "id": adm_id,
        "patientId": data.get("patientId", "PAT-1001"),
        "patientName": data.get("patientName", ""),
        "roomNumber": data.get("roomNumber", ""),
        "bedNumber": data.get("bedNumber", "B-01"),
        "department": data.get("department", "General Medicine"),
        "doctor": data.get("doctor", ""),
        "admissionDate": data.get("admissionDate") or datetime.now().strftime("%Y-%m-%d"),
        "expectedDischarge": data.get("expectedDischarge", ""),
        "actualDischarge": data.get("actualDischarge", ""),
        "status": data.get("status", "Admitted"),
        "wardType": data.get("wardType", "General")
    }

    try:
        created = append_sheet_record("Admissions", new_adm)
        # Auto-update patient status to Inpatient in Patients sheet if patient exists
        if new_adm.get("patientId"):
            update_sheet_record("Patients", new_adm["patientId"], {"status": "Inpatient"})
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to save admission", "details": str(e)}), 500


@app.route("/api/admissions/<adm_id>", methods=["PUT"])
def update_admission_route(adm_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    try:
        updated = update_sheet_record("Admissions", adm_id, data)
        if not updated:
            return jsonify({"error": f"Admission with ID '{adm_id}' not found"}), 404

        # If discharged, update patient record
        if data.get("status") == "Discharged" and updated.get("patientId"):
            update_sheet_record("Patients", updated["patientId"], {"status": "Discharged"})

        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update admission", "details": str(e)}), 500


@app.route("/api/admissions/<adm_id>", methods=["DELETE"])
def delete_admission_route(adm_id):
    try:
        success = delete_sheet_record("Admissions", adm_id)
        if not success:
            return jsonify({"error": f"Admission with ID '{adm_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Admission '{adm_id}' deleted", "id": adm_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete admission", "details": str(e)}), 500


# =========================================================================
# 5. MEDICINES REST API ENDPOINTS
# =========================================================================

@app.route("/api/medicines", methods=["GET"])
def get_medicines():
    try:
        medicines = load_sheet_records("Medicines")
        return jsonify(medicines), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Medicines sheet", "details": str(e)}), 500


@app.route("/api/medicines", methods=["POST"])
def add_medicine():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    if not data.get("name"):
        return jsonify({"error": "Field 'name' is required"}), 400

    med_id = data.get("id") or generate_unique_id("Medicines", "MED-", 500)
    qty = int(data.get("quantity", 50))
    stock_status = "Critical" if qty <= 10 else "Low Stock" if qty <= 25 else "In Stock"

    new_med = {
        "id": med_id,
        "name": data.get("name", "").strip(),
        "category": data.get("category", "General"),
        "quantity": qty,
        "unitPrice": float(data.get("unitPrice", 10.0)),
        "expiryDate": data.get("expiryDate", ""),
        "supplier": data.get("supplier", ""),
        "stockStatus": data.get("stockStatus") or stock_status,
        "dosage": data.get("dosage", ""),
        "batchNumber": data.get("batchNumber", f"LOT-{datetime.now().year}")
    }

    try:
        created = append_sheet_record("Medicines", new_med)
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to add medicine", "details": str(e)}), 500


@app.route("/api/medicines/<med_id>", methods=["PUT"])
def update_medicine_route(med_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    # Auto-adjust stockStatus if quantity changed
    if "quantity" in data:
        qty = int(data["quantity"])
        if "stockStatus" not in data:
            data["stockStatus"] = "Critical" if qty <= 10 else "Low Stock" if qty <= 25 else "In Stock"

    try:
        updated = update_sheet_record("Medicines", med_id, data)
        if not updated:
            return jsonify({"error": f"Medicine with ID '{med_id}' not found"}), 404
        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update medicine", "details": str(e)}), 500


@app.route("/api/medicines/<med_id>", methods=["DELETE"])
def delete_medicine_route(med_id):
    try:
        success = delete_sheet_record("Medicines", med_id)
        if not success:
            return jsonify({"error": f"Medicine with ID '{med_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Medicine '{med_id}' deleted", "id": med_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete medicine", "details": str(e)}), 500


# =========================================================================
# 6. BILLS REST API ENDPOINTS
# =========================================================================

@app.route("/api/bills", methods=["GET"])
def get_bills():
    try:
        bills = load_sheet_records("Bills")
        return jsonify(bills), 200
    except Exception as e:
        return jsonify({"error": "Failed to read Bills sheet", "details": str(e)}), 500


@app.route("/api/bills", methods=["POST"])
def add_bill():
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    if not data.get("patientName"):
        return jsonify({"error": "Field 'patientName' is required"}), 400

    bill_id = data.get("id") or generate_unique_id("Bills", "INV-", 8000)

    consult_fee = float(data.get("consultationFee", 0))
    med_charges = float(data.get("medicineCharges", 0))
    room_charges = float(data.get("roomCharges", 0))
    other_charges = float(data.get("otherCharges", 0))

    # Automatically calculate total amount
    calculated_total = consult_fee + med_charges + room_charges + other_charges
    total_amount = float(data.get("totalAmount")) if data.get("totalAmount") is not None else calculated_total

    new_bill = {
        "id": bill_id,
        "patientId": data.get("patientId", "PAT-1001"),
        "patientName": data.get("patientName", ""),
        "consultationFee": consult_fee,
        "medicineCharges": med_charges,
        "roomCharges": room_charges,
        "otherCharges": other_charges,
        "totalAmount": total_amount,
        "paymentStatus": data.get("paymentStatus", "Pending"),
        "billingDate": data.get("billingDate") or datetime.now().strftime("%Y-%m-%d"),
        "paymentMethod": data.get("paymentMethod", "Credit Card"),
        "insuranceProvider": data.get("insuranceProvider", "")
    }

    try:
        created = append_sheet_record("Bills", new_bill)
        return jsonify(created), 201
    except Exception as e:
        return jsonify({"error": "Failed to generate bill", "details": str(e)}), 500


@app.route("/api/bills/<bill_id>", methods=["PUT"])
def update_bill_route(bill_id):
    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON body provided"}), 400

    # If sub-charges are passed without totalAmount, recalculate totalAmount
    if any(k in data for k in ["consultationFee", "medicineCharges", "roomCharges", "otherCharges"]):
        if "totalAmount" not in data:
            # Load existing record to get missing components
            records = load_sheet_records("Bills")
            existing = next((b for b in records if str(b.get("id")).strip() == str(bill_id).strip()), {})
            c = float(data.get("consultationFee", existing.get("consultationFee", 0)))
            m = float(data.get("medicineCharges", existing.get("medicineCharges", 0)))
            r = float(data.get("roomCharges", existing.get("roomCharges", 0)))
            o = float(data.get("otherCharges", existing.get("otherCharges", 0)))
            data["totalAmount"] = c + m + r + o

    try:
        updated = update_sheet_record("Bills", bill_id, data)
        if not updated:
            return jsonify({"error": f"Bill with ID '{bill_id}' not found"}), 404
        return jsonify(updated), 200
    except Exception as e:
        return jsonify({"error": "Failed to update bill", "details": str(e)}), 500


@app.route("/api/bills/<bill_id>", methods=["DELETE"])
def delete_bill_route(bill_id):
    try:
        success = delete_sheet_record("Bills", bill_id)
        if not success:
            return jsonify({"error": f"Bill with ID '{bill_id}' not found"}), 404
        return jsonify({"success": True, "message": f"Bill '{bill_id}' deleted", "id": bill_id}), 200
    except Exception as e:
        return jsonify({"error": "Failed to delete bill", "details": str(e)}), 500


# -------------------------------------------------------------------------
# APPLICATION ENTRYPOINT
# -------------------------------------------------------------------------
if __name__ == "__main__":
    print(f"============================================================")
    print(f" AuraCare Hospital Management System - REST Backend")
    print(f" Database Workbook: {EXCEL_FILE}")
    print(f" Server Port: 5000")
    print(f" CORS: Enabled for all origins (*)")
    print(f"============================================================")
    app.run(host="0.0.0.0", port=5000, debug=True)
